# -*- coding: utf-8 -*-
"""
GeoSalud Cotopaxi - Supplementary Material generator (run in Google Colab or locally).

INPUT  weekly_panel.csv  (one row per canton-year-week; the SAME final data used for Table 2)
   columns: canton, year, week, so2_raw, discharges
            (so2_raw = weekly mean TROPOMI SO2_column_number_density, mol/m2, NaN when missing, NOT replaced)
   optional: wind_top5.csv  columns: year, week, direction_toward_deg   (ERA5-Land 10 m wind, five highest-SO2 weeks)
   optional: Supp_Figure_2023_SO2_descriptive.png, Supp_scatterplots_SO2_vs_discharges.png (Figures S1, S2)

OUTPUT Supplementary_Material.docx + one CSV per table.
   Table S2  leave-one-out         Table S3  autocorrelation        Table S4  NaN / MAD counts
   Table S5a cleaned correlations + Fisher test    Table S5b cleaned quasi-Poisson (4 core cantons)
   Table S6  wind check            Table S7  IQR + full quasi-Poisson output + GLM residual autocorrelation
Nothing is invented: every number is computed from weekly_panel.csv.
"""
import numpy as np, pandas as pd, os
from scipy import stats
import statsmodels.api as sm
from statsmodels.stats.stattools import durbin_watson
from docx import Document
from docx.shared import Inches

CORE = ['Latacunga', 'Pujilí', 'Saquisilí', 'Salcedo']
THRESH = 0.0002          # mol/m2, empirical screen (Methods)
df = pd.read_csv('weekly_panel.csv')
df = df[df.year.isin([2022, 2024])].sort_values(['canton', 'year', 'week']).reset_index(drop=True)

def fisher(r1, n1, r2, n2):
    z = (np.arctanh(r1) - np.arctanh(r2)) / np.sqrt(1/(n1-3) + 1/(n2-3))
    return z, 2*(1 - stats.norm.cdf(abs(z)))

def pear(x, y):
    m = x.notna() & y.notna()
    r, p = stats.pearsonr(x[m], y[m]); return r, p, int(m.sum())

# ---------- cleaned series (secondary): MAD replacement, 2022 weeks 40-52 exempt, NaN -> median
def clean(g):
    s = g.so2_raw.copy(); med = s.median(); mad = (s - med).abs().median()
    flag = (s > med + 2.5*mad) & (s > THRESH)
    if g.name[1] == 2022: flag &= ~g.week.between(40, 52)
    out = s.copy(); out[flag] = med; out[s.isna()] = med
    return pd.DataFrame({'so2_clean': out, 'mad_flag': flag, 'nan_flag': s.isna()}, index=g.index)
df = df.join(df.groupby(['canton', 'year'], group_keys=False).apply(clean))

# ---------- S4: counts
S4 = df.groupby(['canton', 'year']).agg(n_weeks=('week', 'size'), n_NaN=('nan_flag', 'sum'), n_MAD_flagged=('mad_flag', 'sum')).reset_index()
S4['pct_modified'] = ((S4.n_NaN + S4.n_MAD_flagged) / S4.n_weeks * 100).round(1)

# ---------- S2 leave-one-out, S3 autocorrelation (raw series)
S2, S3 = [], []
for (c, y), g in df.groupby(['canton', 'year']):
    m = g.so2_raw.notna() & g.discharges.notna(); h = g[m]
    r, p = stats.pearsonr(h.so2_raw, h.discharges)
    k = h.so2_raw.idxmax(); h2 = h.drop(k)
    r2, p2 = stats.pearsonr(h2.so2_raw, h2.discharges)
    S2.append(dict(canton=c, year=y, n=len(h), r=round(r, 3), p=round(p, 4), dropped_week=int(h.loc[k, 'week']), r_LOO=round(r2, 3), p_LOO=round(p2, 4)))
    ols = sm.OLS(h.discharges.values, sm.add_constant(h.so2_raw.values)).fit(); e = ols.resid
    S3.append(dict(canton=c, year=y, n=len(h), durbin_watson=round(durbin_watson(e), 2), lag1_resid_autocorr=round(np.corrcoef(e[:-1], e[1:])[0, 1], 3)))
S2, S3 = pd.DataFrame(S2), pd.DataFrame(S3)

# ---------- S5a cleaned correlations + Fisher (cleaned) ; Table 2 raw also given for comparison
rows = []
for c, gc in df.groupby('canton'):
    d = {}
    for y, g in gc.groupby('year'):
        d[y] = (pear(g.so2_clean, g.discharges), pear(g.so2_raw, g.discharges))
    (rc22, pc22, n22), (rr22, pr22, _) = d[2022]; (rc24, pc24, n24), (rr24, pr24, _) = d[2024]
    zc, pzc = fisher(rc22, n22, rc24, n24); zr, pzr = fisher(rr22, d[2022][1][2], rr24, d[2024][1][2])
    rows.append(dict(canton=c, r2022_raw=round(rr22, 3), p2022_raw=round(pr22, 4), r2024_raw=round(rr24, 3), p2024_raw=round(pr24, 4), fisher_p_raw=round(pzr, 4),
                     r2022_clean=round(rc22, 3), p2022_clean=round(pc22, 4), r2024_clean=round(rc24, 3), p2024_clean=round(pc24, 4), fisher_p_clean=round(pzc, 4)))
S5a = pd.DataFrame(rows)

# ---------- quasi-Poisson models: discharges ~ pred (+ week + week^2 when adjusted); same method as the Colab run
def qp_models(series_col, adjusted=True):
    out = []
    for c in CORE:
        for y in (2022, 2024):
            g = df[(df.canton == c) & (df.year == y)].sort_values('week').reset_index(drop=True)
            s = g[series_col]
            lags = {f'lag{k}': s.shift(k) for k in range(4)}
            lags['lag_avg03'] = s.rolling(4, min_periods=4).mean()
            for name, v in lags.items():
                d = pd.DataFrame({'y': g.discharges, 'x': v, 'w': g.week.astype(float)}).dropna()
                iqr = d.x.quantile(.75) - d.x.quantile(.25)
                cols = ['x'] + (['w', 'w2'] if adjusted else [])
                d['w2'] = d.w ** 2
                X = sm.add_constant(d[cols])
                res = sm.GLM(d.y, X, family=sm.families.Poisson()).fit()
                disp = res.pearson_chi2 / res.df_resid
                b, se = res.params['x'], res.bse['x'] * np.sqrt(max(disp, 1.0))
                p = 2 * (1 - stats.norm.cdf(abs(b / se))); e = res.resid_pearson.values
                out.append(dict(canton=c, year=y, lag=name, n=len(d), IQR_mol_m2=round(iqr, 6), beta_per_unit=round(b, 1),
                                ER_pct=round((np.exp(b*iqr)-1)*100, 1), CI95_low=round((np.exp((b-1.96*se)*iqr)-1)*100, 1),
                                CI95_high=round((np.exp((b+1.96*se)*iqr)-1)*100, 1), p_quasi=round(p, 4), dispersion=round(disp, 2),
                                DW_resid=round(durbin_watson(e), 2), lag1_resid_autocorr=round(np.corrcoef(e[:-1], e[1:])[0, 1], 3)))
    return pd.DataFrame(out)
S7 = qp_models('so2_raw')            # primary: raw SO2, adjusted for week and week^2
S5b = qp_models('so2_clean')         # secondary: MAD-cleaned SO2, adjusted

# ---------- S6 wind check (optional input)
S6 = pd.read_csv('wind_top5.csv') if os.path.exists('wind_top5.csv') else None

tables = {'S2': (S2, 'Leave-one-out sensitivity of the primary (raw-data) Pearson correlations: the single highest-SO₂ week of each canton and year was excluded.'),
          'S3': (S3, 'Temporal-autocorrelation diagnostic (discharges ~ SO₂ simple regression) for each canton and year.'),
          'S4': (S4, 'Number of missing (NaN) weeks and MAD-flagged outliers by canton and year (secondary, MAD-cleaned series).'),
          'S5a': (S5a, 'Secondary analysis: Pearson correlations and Fisher r-to-z p-values with the MAD-cleaned series, next to the primary raw-data values.'),
          'S5b': (S5b, 'Secondary analysis: quasi-Poisson excess risk (ER%) with the MAD-cleaned series, adjusted for epidemiological week and its square, four core cantons.'),
          'S7': (S7, 'Primary analysis: IQR, regression output, dispersion and residual autocorrelation for every quasi-Poisson model (raw SO₂, adjusted for epidemiological week and its square).')}
if S6 is not None: tables['S6'] = (S6, 'Post hoc wind check: ERA5-Land 10 m wind direction (toward) at the crater for the five highest-SO₂ weeks of each year.')
for k, (t, _) in tables.items(): t.to_csv(f'Table_{k}.csv', index=False)

# ---------- assemble Supplementary_Material.docx
doc = Document(); doc.add_heading('Supplementary Material', 0)
doc.add_paragraph('Spatiotemporal Association Between Cotopaxi Volcano SO₂ Emissions and Respiratory Hospital Discharges in Cotopaxi Province, Ecuador, 2022 and 2024.')
for fig, f, cap in [('S1', 'Supp_Figure_2023_SO2_descriptive.png', 'Weekly SO₂ column density, 2022 vs 2023 vs 2024 (descriptive context for the exclusion of 2023).'),
                    ('S2', 'Supp_scatterplots_SO2_vs_discharges.png', 'Weekly SO₂ column density vs respiratory discharges, all cantons, 2022 and 2024.')]:
    if os.path.exists(f):
        doc.add_picture(f, width=Inches(5.8 if fig == 'S1' else 3.6)); doc.add_paragraph(f'Supplementary Figure {fig}. {cap}')
for k in ['S2', 'S3', 'S4', 'S5a', 'S5b', 'S6', 'S7']:
    if k not in tables: continue
    t, cap = tables[k]; doc.add_paragraph(f'Supplementary Table {k}. {cap}')
    tb = doc.add_table(rows=1, cols=len(t.columns)); tb.style = 'Light Grid Accent 1'
    for i, h in enumerate(t.columns): tb.rows[0].cells[i].text = str(h)
    for _, r in t.iterrows():
        cells = tb.add_row().cells
        for i, v in enumerate(r): cells[i].text = '' if pd.isna(v) else str(v)
doc.save('Supplementary_Material.docx'); print('done:', list(tables))
